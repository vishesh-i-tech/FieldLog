import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def create_document():
    doc = Document()

    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Title
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    run_title = title_p.add_run("FieldLog — Enterprise Daily Operations & Field Reporting System")
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(37, 99, 235) # #2563EB

    # Subtitle
    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_after = Pt(16)
    run_sub = sub_p.add_run("Complete System Architecture, Multi-Tenant Workflows, and Technical Specification")
    run_sub.font.name = "Calibri"
    run_sub.font.size = Pt(13)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(100, 116, 139)

    # Metadata Table
    table = doc.add_table(rows=4, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    meta_data = [
        ("Project Name", "FieldLog Monorepo (Web, Android, iOS & API)"),
        ("Version", "v1.2.0 (Production Architecture)"),
        ("Repository", "https://github.com/vishesh-i-tech/FieldLog.git"),
        ("Last Updated", "October 2026"),
    ]

    for i, (k, v) in enumerate(meta_data):
        row = table.rows[i]
        c0 = row.cells[0]
        c1 = row.cells[1]
        c0.width = Inches(2.2)
        c1.width = Inches(4.5)
        set_cell_background(c0, "F1F5F9")
        set_cell_background(c1, "FFFFFF")
        
        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_after = Pt(2)
        r0 = p0.add_run(k)
        r0.font.bold = True
        r0.font.size = Pt(10)
        
        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_after = Pt(2)
        r1 = p1.add_run(v)
        r1.font.size = Pt(10)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # 1. Executive Summary
    h1 = doc.add_heading("1. Executive Summary", level=1)
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(6)

    doc.add_paragraph(
        "FieldLog is an end-to-end multi-tenant operations platform designed to capture, verify, and monitor "
        "daily field progress across distributed work places (construction sites, manufacturing units, warehouses, retail hubs). "
        "The system enforces strict operational integrity through real-device GPS geotagging, on-device canvas watermarking, "
        "audio voice notes, standardized issue tracking, and instant role-based access for Owners and Field Officers."
    )

    # 2. Key Architecture Pillars
    h1 = doc.add_heading("2. Architectural Overview", level=1)
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(6)

    pillars = [
        ("Universal Mobile & Web Frontend", "Built on React Native and Expo SDK 52 with Expo Router. Ships as a high-performance Web App and native Android/iOS mobile builds from a unified code base."),
        ("High-Performance Backend", "Fastify v5 framework with ultra-low latency, Zod schema validation, and JSON Web Token (JWT) multi-tenant authentication."),
        ("Type-Safe Database Layer", "Drizzle ORM mapping relational schemas with SQLite for rapid local development and PostgreSQL for enterprise cloud deployments."),
        ("Tamper-Proof Geotagging Engine", "Extracts real latitude and longitude from photo EXIF APP1 binary segments or device GPS, stamping verifiable watermarks directly into the image canvas prior to storage."),
        ("Multi-Tenant Isolation", "Organizations are partitioned with role-based policies separating Owners (broad oversight) from Field Officers (strictly limited to assigned work places)."),
    ]

    for title, desc in pillars:
        p = doc.add_paragraph(style='List Bullet')
        r_t = p.add_run(f"{title}: ")
        r_t.font.bold = True
        p.add_run(desc)

    # 3. Roles and Workflow
    h1 = doc.add_heading("3. Roles & Operational Workflow", level=1)
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(6)

    doc.add_heading("3.1 Owner / Admin Dashboard", level=2)
    doc.add_paragraph(
        "The Owner dashboard (/owner) provides a single-pane-of-glass overview of all organizational sites:\n"
        "• Real-time submission counter (Full Day, Half Day, Work Stopped, Pending).\n"
        "• Work Place Management: Create new sites, assign locations, and assign Field Officers.\n"
        "• Field Officer Management: Invite supervisors with mobile phone numbers and 4-digit PINs.\n"
        "• Field Officer Assignment: Every site card displays assigned officer names and phone numbers with quick 1-click reassignment.\n"
        "• Detail Inspection: Drill down into individual submissions to view full-resolution geotagged photos, listen to audio notes, inspect reported blockers, and mark records as Reviewed."
    )

    doc.add_heading("3.2 Field Officer Portal", level=2)
    doc.add_paragraph(
        "The Field Officer interface (/record) provides a streamlined, single-screen reporting workflow:\n"
        "• Quick Login: Passwordless phone + PIN authentication.\n"
        "• Assigned Site Selector: Displays only sites explicitly assigned by the owner.\n"
        "• Attendance Status: 1-tap toggling between Full Day, Half Day, or Work Stopped.\n"
        "• Camera & Photo Attachment: Live camera capture or device gallery upload with automatic EXIF GPS extraction and canvas watermarking.\n"
        "• Voice Note: One-tap voice recording (up to 60 seconds) for hands-free progress briefings.\n"
        "• Issue Tracker: Select pre-configured issue chips (e.g., Material Delay, Weather) or enter custom descriptions."
    )

    # 4. Data Models & Schemas
    h1 = doc.add_heading("4. Database Schema Specification", level=1)
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(6)

    schemas = [
        ("organizations", "id (UUID), name, plan, createdAt, updatedAt"),
        ("users", "id (UUID), organizationId, role ('owner' | 'supervisor'), phone, name, pinHash, language, isActive"),
        ("sites", "id (UUID), organizationId, name, locationText, latitude, longitude, isArchived, createdAt"),
        ("site_assignments", "id (UUID), siteId (FK), userId (FK), assignedAt"),
        ("daily_records", "id (UUID), organizationId, siteId, recordDate, status, capturedAt, missingPhoto, missingVoice, submittedBy"),
        ("record_issues", "id (UUID), recordId (FK), issueTypeId (FK)"),
        ("record_media", "id (UUID), recordId (FK), kind ('photo' | 'audio'), storageKey, mimeType, latitude, longitude, durationSeconds"),
    ]

    t_schema = doc.add_table(rows=len(schemas)+1, cols=2)
    t_schema.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t_schema.rows[0]
    hdr.cells[0].paragraphs[0].add_run("Table Name").font.bold = True
    hdr.cells[1].paragraphs[0].add_run("Columns & Relationships").font.bold = True
    set_cell_background(hdr.cells[0], "E2E8F0")
    set_cell_background(hdr.cells[1], "E2E8F0")

    for i, (t_name, cols) in enumerate(schemas):
        r = t_schema.rows[i+1]
        r.cells[0].width = Inches(2.0)
        r.cells[1].width = Inches(4.7)
        r.cells[0].paragraphs[0].add_run(t_name).font.bold = True
        r.cells[1].paragraphs[0].add_run(cols)
        bg = "F8FAFC" if i % 2 == 0 else "FFFFFF"
        set_cell_background(r.cells[0], bg)
        set_cell_background(r.cells[1], bg)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # 5. Core API Endpoints
    h1 = doc.add_heading("5. API Route Specification", level=1)
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(6)

    api_routes = [
        ("POST /api/v1/auth/owner/signup", "Owner registration with organization creation."),
        ("POST /api/v1/auth/supervisor/login", "Field Officer authentication via phone & 4-digit PIN."),
        ("GET /api/v1/sites", "Lists all organization sites (Owner) or assigned sites (Supervisor)."),
        ("POST /api/v1/sites", "Creates a new work place (Owner only)."),
        ("POST /api/v1/sites/:id/assign", "Assigns one or more Field Officers to a work place."),
        ("DELETE /api/v1/sites/:id/assign/:userId", "Unassigns a Field Officer from a work place."),
        ("GET /api/v1/supervisors", "Retrieves all registered Field Officers for the organization."),
        ("POST /api/v1/supervisors", "Adds a new Field Officer with credentials."),
        ("GET /api/v1/records/today", "Aggregates today's submission statuses across all sites."),
        ("POST /api/v1/records", "Submits a daily operational record with photos, voice, and issues."),
        ("GET /api/v1/records/:id", "Fetches complete record inspection details and media assets."),
        ("PUT /api/v1/dev-storage/upload", "Local binary and base64 media asset persistence."),
        ("GET /api/v1/dev-storage/files/:key", "Serves media files with verified MIME types."),
    ]

    for ep, desc in api_routes:
        p = doc.add_paragraph(style='List Bullet')
        r_ep = p.add_run(f"{ep}: ")
        r_ep.font.bold = True
        r_ep.font.name = "Consolas"
        r_ep.font.size = Pt(9.5)
        p.add_run(desc)

    # 6. Local Setup & Execution Guide
    h1 = doc.add_heading("6. Setup & Execution Guide", level=1)
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(6)

    doc.add_paragraph("To run the FieldLog development stack locally:")

    steps = [
        ("Install Dependencies", "Run `npm install` from the repository root to install backend and mobile dependencies."),
        ("Database Migrations & Seed", "Run `npm run db:push` followed by `npm run seed:dev` to seed demo sites, officers, and initial records."),
        ("Launch Fastify Backend", "Run `npm run dev:backend` (runs on http://localhost:4000)."),
        ("Launch Expo Metro Web", "Run `npm run dev:web` (runs on http://localhost:8081)."),
        ("Run Backend Test Suites", "Run `npm run test:backend` (validates all 6 multi-tenant lifecycle test suites)."),
    ]

    for title, cmd in steps:
        p = doc.add_paragraph(style='List Bullet')
        p.add_run(f"{title}: ").font.bold = True
        p.add_run(cmd)

    output_path = os.path.abspath("docs/FieldLog_Project_Documentation.docx")
    doc.save(output_path)
    print(f"Successfully generated documentation: {output_path}")

if __name__ == "__main__":
    create_document()
